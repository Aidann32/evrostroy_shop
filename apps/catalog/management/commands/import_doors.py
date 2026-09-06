import os
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.files import File
from django.core.management.base import BaseCommand
from django.utils.text import slugify
from openpyxl import load_workbook
from unidecode import unidecode

from apps.catalog.models import (
    Attribute,
    Brand,
    Category,
    Product,
    ProductAttribute,
)


# Базовые колонки Excel (не считаются характеристиками)
BASE_COLUMNS = {
    'id',
    'title',
    'description',
    'brand',
    'category',
    'price',
    'фулл прайс',  # цена комплекта для дверей
}


def make_slug(text):
    return slugify(unidecode(str(text))) or 'item'


class Command(BaseCommand):
    help = 'Импорт товаров из Excel (краски, двери и т.д.)'

    def add_arguments(self, parser):
        parser.add_argument(
            'excel_path',
            type=str,
            help='Путь к Excel, например: /app/data/двери.xlsx',
        )
        parser.add_argument(
            '--images-dir',
            type=str,
            default='',
            help='Папка с картинками (имя файла = id)',
        )
        parser.add_argument(
            '--dry-run',
            action='store_true',
            help='Только показать, что будет импортировано',
        )

    def handle(self, *args, **options):
        excel_path = options['excel_path']
        images_dir = options['images_dir']
        dry_run = options['dry_run']

        if not os.path.exists(excel_path):
            self.stderr.write(self.style.ERROR(f'Файл не найден: {excel_path}'))
            return

        wb = load_workbook(excel_path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            self.stderr.write(self.style.ERROR('Excel пустой'))
            return

        header = [str(c).strip().lower() if c is not None else '' for c in rows[0]]
        col = {name: idx for idx, name in enumerate(header)}

        # price или «фулл прайс»
        price_key = None
        if 'price' in col:
            price_key = 'price'
        elif 'фулл прайс' in col:
            price_key = 'фулл прайс'

        if 'id' not in col or 'title' not in col or price_key is None:
            self.stderr.write(self.style.ERROR(
                'Нужны колонки: id, title и price (или «фулл прайс»)'
            ))
            return

        # Все колонки, кроме базовых → характеристики
        attr_columns = [
            name for name in header
            if name and name not in BASE_COLUMNS
        ]

        created = updated = skipped = 0

        for row in rows[1:]:
            if not row or all(c is None or str(c).strip() == '' for c in row):
                continue

            raw_id = row[col['id']]
            title = str(row[col['title']] or '').strip()
            description = ''
            if 'description' in col:
                description = str(row[col['description']] or '').strip()

            brand_name = 'Без бренда'
            if 'brand' in col:
                brand_name = str(row[col['brand']] or '').strip() or 'Без бренда'

            category_name = ''
            if 'category' in col:
                category_name = str(row[col['category']] or '').strip()

            price_raw = row[col[price_key]]

            if not title:
                skipped += 1
                continue

            # Уникальное имя: для дверей title одинаковый
            if description:
                name = f'{title}: {description}'
            else:
                name = title

            try:
                price = Decimal(str(price_raw).replace(' ', '').replace(',', '.'))
            except (InvalidOperation, TypeError, ValueError):
                self.stderr.write(self.style.WARNING(f'Пропуск (цена): {name}'))
                skipped += 1
                continue

            product_id_key = str(raw_id).strip() if raw_id is not None else ''

            # Значения характеристик из доп. столбцов
            attributes_data = {}
            for attr_name in attr_columns:
                idx = col[attr_name]
                if idx >= len(row):
                    continue
                val = row[idx]
                if val is None or str(val).strip() == '':
                    continue
                attributes_data[attr_name.strip()] = str(val).strip()

            if dry_run:
                attrs_str = ', '.join(f'{k}={v}' for k, v in attributes_data.items())
                self.stdout.write(
                    f'[DRY] {product_id_key} | {name} | {brand_name} | '
                    f'{category_name} | {price} | {attrs_str}'
                )
                created += 1
                continue

            brand, _ = Brand.objects.get_or_create(
                name=brand_name,
                defaults={'slug': self._unique_slug(Brand, brand_name)},
            )

            category = None
            if category_name:
                category, _ = Category.objects.get_or_create(
                    name=category_name,
                    defaults={'slug': self._unique_slug(Category, category_name)},
                )

            product = Product.objects.filter(name=name, brand=brand).first()

            if product:
                product.description = description
                product.category = category
                product.price = price
                product.is_active = True
                if not product.slug:
                    product.slug = self._unique_slug(Product, name, instance=product)
                product.save()
                updated += 1
                action = 'обновлён'
            else:
                product = Product(
                    name=name,
                    brand=brand,
                    category=category,
                    description=description,
                    price=price,
                    is_active=True,
                )
                product.slug = self._unique_slug(Product, name)
                product.save()
                created += 1
                action = 'создан'

            # Характеристики: полотно, коробка, наличник и т.д.
            for attr_name, attr_value in attributes_data.items():
                attribute, _ = Attribute.objects.get_or_create(
                    name=attr_name.capitalize() if attr_name.islower() else attr_name,
                    defaults={'slug': self._unique_slug(Attribute, attr_name)},
                )
                # Красивые названия: полотно → Полотно
                if attribute.name != attr_name and attr_name in ('полотно', 'коробка', 'наличник'):
                    pretty = attr_name.capitalize()
                    if attribute.name != pretty:
                        attribute.name = pretty
                        attribute.save(update_fields=['name'])

                ProductAttribute.objects.update_or_create(
                    product=product,
                    attribute=attribute,
                    defaults={'value': attr_value},
                )

            # Картинка
            if images_dir and product_id_key:
                image_path = self._find_image(images_dir, product_id_key)
                if image_path:
                    with open(image_path, 'rb') as f:
                        product.image.save(os.path.basename(image_path), File(f), save=True)
                    self.stdout.write(self.style.SUCCESS(f'{action}: {name} + image'))
                else:
                    self.stdout.write(self.style.WARNING(
                        f'{action}: {name} (нет картинки id={product_id_key})'
                    ))
            else:
                self.stdout.write(self.style.SUCCESS(f'{action}: {name}'))

        wb.close()
        self.stdout.write('')
        self.stdout.write(self.style.SUCCESS(
            f'Готово. Создано: {created}, обновлено: {updated}, пропущено: {skipped}'
        ))

    def _unique_slug(self, model, text, instance=None):
        base = make_slug(text)[:240]
        slug = base
        n = 1
        qs = model.objects.all()
        if instance and instance.pk:
            qs = qs.exclude(pk=instance.pk)
        while qs.filter(slug=slug).exists():
            slug = f'{base}-{n}'
            n += 1
        return slug[:255]

    def _find_image(self, images_dir, product_id):
        images_dir = Path(images_dir)
        if not images_dir.exists():
            return None

        extensions = ['.jpg', '.jpeg', '.png', '.webp', '.JPG', '.JPEG', '.PNG', '.WEBP']
        for ext in extensions:
            path = images_dir / f'{product_id}{ext}'
            if path.is_file():
                return str(path)

        for path in images_dir.iterdir():
            if path.is_file() and path.stem == str(product_id):
                return str(path)
        return None